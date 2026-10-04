"""Downloadable presentation (.pptx) and document (.docx) versions of a report.

Both are built from the finished report deterministically - no model call, no extra
cost - so they say exactly what the report says, with the same citation handles and
the same methodology and limitations. Charts in the deck are native PowerPoint charts
(editable, not pictures), so they can be restyled in PowerPoint.
"""

from __future__ import annotations

import io
from typing import Any

from .models import DESIGN_LABELS, TIER_MEANING, Report

NAVY = (0x0B, 0x14, 0x26)
INK = (0x1F, 0x29, 0x37)
MUTED = (0x6B, 0x72, 0x80)
BLUE = (0x25, 0x63, 0xEB)
VIOLET = (0x7C, 0x3A, 0xED)
GREEN = (0x05, 0x96, 0x69)
AMBER = (0xD9, 0x77, 0x06)
RED = (0xDC, 0x26, 0x26)
STRENGTH_COLOR = {"strong": GREEN, "moderate": AMBER, "limited": (0xEA, 0x58, 0x0C), "preliminary": RED}


def _clip(text: str, limit: int) -> str:
    text = " ".join((text or "").split())
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _cites(handles: list[str]) -> str:
    return " ".join(f"[{h}]" for h in handles) if handles else "no resolvable citation"


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #

def to_pptx(report: Report) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt

    rgb = lambda c: RGBColor(*c)  # noqa: E731
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    blank = deck.slide_layouts[6]
    s, m = report.synthesis, report.metrics

    def text(slide, left, top, width, height, value, *, size=16, bold=False, color=INK, align=None):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = True
        lines = value if isinstance(value, list) else [value]
        for i, line in enumerate(lines):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = line
            run.font.size, run.font.bold, run.font.color.rgb = Pt(size), bold, rgb(color)
            run.font.name = "Segoe UI"
            if align:
                paragraph.alignment = align
            paragraph.space_after = Pt(6)
        return box

    def rect(slide, left, top, width, height, fill, shape=MSO_SHAPE.RECTANGLE):
        box = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
        box.fill.solid()
        box.fill.fore_color.rgb = rgb(fill)
        box.line.fill.background()
        return box

    def page(title: str, kicker: str = ""):
        slide = deck.slides.add_slide(blank)
        rect(slide, 0, 0, 13.333, 0.12, BLUE)
        rect(slide, 0, 0, 4.0, 0.12, VIOLET)
        if kicker:
            text(slide, 0.6, 0.35, 12, 0.4, kicker.upper(), size=11, bold=True, color=MUTED)
        text(slide, 0.6, 0.65, 12.1, 0.9, title, size=28, bold=True, color=NAVY)
        text(slide, 0.6, 7.0, 12.1, 0.35,
             "DrugScope research brief · research intelligence, not medical advice · citations refer to the Sources slide",
             size=9, color=MUTED)
        return slide

    def bars(slide, categories, values, left, top, width, height, color=BLUE, horizontal=False):
        data = CategoryChartData()
        data.categories = categories
        data.add_series("Count", values)
        kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
        chart = slide.shapes.add_chart(kind, Inches(left), Inches(top), Inches(width), Inches(height), data).chart
        chart.has_legend = False
        plot = chart.plots[0]
        plot.gap_width = 60
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(11)
        plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        plot.series[0].format.fill.solid()
        plot.series[0].format.fill.fore_color.rgb = rgb(color)
        chart.category_axis.tick_labels.font.size = Pt(11)
        axis = chart.value_axis
        axis.has_major_gridlines = False
        axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        axis.major_tick_mark = XL_TICK_MARK.NONE
        axis.format.line.fill.background()
        if horizontal:
            chart.category_axis.reverse_order = True
        return chart

    # 1. Title
    slide = deck.slides.add_slide(blank)
    rect(slide, 0, 0, 13.333, 7.5, NAVY)
    rect(slide, 0, 0, 0.35, 7.5, BLUE)
    rect(slide, 0.35, 0, 0.15, 7.5, VIOLET)
    text(slide, 1.0, 1.0, 11, 0.5, "DRUGSCOPE RESEARCH BRIEF", size=14, bold=True, color=(0x93, 0xC5, 0xFD))
    text(slide, 1.0, 1.7, 11.3, 2.8, _clip(report.query, 160), size=36, bold=True, color=(0xFF, 0xFF, 0xFF))
    text(slide, 1.0, 4.6, 11.3, 1.0, _clip(report.plan.interpretation, 260), size=16, color=(0xCB, 0xD5, 0xE1))
    text(slide, 1.0, 6.2, 11.3, 0.5,
         f"{report.generated_at[:10]} · {report.depth} depth · {len(report.records)} sources · "
         f"Evidence Quality {m['quality']['score']}/100",
         size=13, color=(0x94, 0xA3, 0xB8))

    # 2. Bottom line
    slide = page("Bottom line", f"{s.summary.confidence} confidence")
    rect(slide, 0.6, 1.75, 0.08, 2.4, BLUE)
    text(slide, 0.9, 1.7, 11.8, 2.6, _clip(s.summary.verdict, 420), size=24, bold=True, color=NAVY)
    text(slide, 0.9, 4.5, 11.8, 2.2, ["Why this confidence:", _clip(s.summary.confidence_rationale, 600)],
         size=15, color=INK)

    # 3. Evidence at a glance
    slide = page("Evidence at a glance", "computed from the retrieved corpus")
    tiles = [
        ("Evidence quality", f"{m['quality']['score']}/100", BLUE),
        ("Consensus", f"{m['consensus']['score']}/100", GREEN),
        ("Development maturity", f"{s.summary.maturity}/100", VIOLET),
        ("Sources analysed", str(m["totals"]["records"]), AMBER),
    ]
    for i, (label, value, color) in enumerate(tiles):
        left = 0.6 + i * 3.05
        rect(slide, left, 1.75, 2.85, 1.45, (0xF3, 0xF6, 0xFB), MSO_SHAPE.ROUNDED_RECTANGLE)
        rect(slide, left, 1.75, 2.85, 0.07, color)
        text(slide, left + 0.2, 1.9, 2.5, 0.4, label.upper(), size=10, bold=True, color=MUTED)
        text(slide, left + 0.2, 2.25, 2.5, 0.8, value, size=30, bold=True, color=NAVY)
    tiers = m["literature"]["tier_mix"]
    if any(r["count"] for r in tiers):
        text(slide, 0.6, 3.45, 6, 0.4, "Papers by evidence tier", size=13, bold=True, color=NAVY)
        bars(slide, [r["tier"] for r in tiers], [r["count"] for r in tiers], 0.6, 3.8, 6.0, 3.0)
    designs = [r for r in m["literature"]["design_mix"] if r["count"]][:6]
    if designs:
        text(slide, 6.9, 3.45, 6, 0.4, "Study designs", size=13, bold=True, color=NAVY)
        bars(slide, [r["design"] for r in designs], [r["count"] for r in designs], 6.9, 3.8, 5.8, 3.0,
             color=VIOLET, horizontal=True)

    # 4. Key findings, two per slide so long headlines never collide with their detail
    def card(slide, top, title, body, foot, color, height=2.45):
        rect(slide, 0.6, top, 0.08, height - 0.15, color)
        title_lines = 1 + len(title) // 95
        text(slide, 0.85, top - 0.05, 11.8, 0.45 * title_lines, title, size=17, bold=True, color=NAVY)
        body_top = top + 0.1 + 0.42 * title_lines
        text(slide, 0.85, body_top, 11.8, height - 0.55 - 0.42 * title_lines, body, size=13, color=INK)
        text(slide, 0.85, top + height - 0.5, 11.8, 0.35, foot, size=11, bold=True, color=color)

    findings = s.key_findings
    for start in range(0, len(findings), 2):
        chunk = findings[start:start + 2]
        slide = page("Key findings" + (f" ({start // 2 + 1} of {(len(findings) + 1) // 2})" if len(findings) > 2 else ""),
                     "ordered by how much they should change a decision")
        for row, finding in enumerate(chunk):
            color = STRENGTH_COLOR.get(finding.strength, AMBER)
            card(slide, 1.7 + row * 2.65, f"{start + row + 1}. {_clip(finding.headline, 180)}",
                 _clip(finding.detail, 520), f"Strength: {finding.strength} · Evidence: {_cites(finding.citations)}",
                 color)

    # 5. Trials
    trials = m.get("trials") or {}
    if trials.get("count"):
        slide = page("Clinical trial landscape", f"{trials['count']} trials retrieved")
        facts = [
            f"Highest phase: {trials.get('highest_phase', 'n/a')}",
            f"Active: {trials.get('active', 0)} · completed: {trials.get('completed', 0)} · "
            f"stopped: {trials.get('stopped', 0)}",
            f"Participants: {trials.get('total_enrollment', 0):,} (median {trials.get('median_enrollment', 0):,})",
            f"With posted results: {trials.get('with_results', 0)}",
            f"Industry-sponsored: {trials.get('industry_share', 0)}%",
        ]
        text(slide, 0.6, 1.75, 5.0, 4.5, facts, size=15, color=INK)
        phases = [r for r in trials.get("phase_mix", []) if r["count"]]
        if phases:
            bars(slide, [r["phase"] for r in phases], [r["count"] for r in phases], 5.9, 1.7, 6.9, 5.0, color=GREEN)

    # 6. Comparison
    comparison = s.comparison
    if comparison.entities and comparison.rows:
        slide = page("Head-to-head comparison", _clip(comparison.bottom_line, 140))
        entities = comparison.entities[:4]
        rows = comparison.rows[:7]
        table = slide.shapes.add_table(len(rows) + 1, len(entities) + 2, Inches(0.6), Inches(1.7),
                                       Inches(12.1), Inches(0.5 + 0.6 * len(rows))).table
        headers = ["Criterion", *entities, "Verdict"]
        for col, header in enumerate(headers):
            cell = table.cell(0, col)
            cell.text = header
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
            cell.text_frame.paragraphs[0].runs[0].font.bold = True
        for r, row in enumerate(rows, 1):
            by_entity = {c.entity: c for c in row.cells}
            values = [row.criterion] + [
                (_clip(by_entity[e].value, 90) + (" " + _cites(by_entity[e].citations) if by_entity[e].citations else ""))
                if e in by_entity else "Not established" for e in entities
            ] + [_clip(row.verdict or "-", 80)]
            for col, value in enumerate(values):
                cell = table.cell(r, col)
                cell.text = value
                for paragraph in cell.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(10)

    # 7. Safety
    severity_color = {"critical": RED, "serious": (0xEA, 0x58, 0x0C), "moderate": AMBER, "mild": GREEN}
    signals = s.safety_signals
    for start in range(0, len(signals), 3):
        slide = page("Safety signals", "from labels, trials and post-market reporting")
        for row, x in enumerate(signals[start:start + 3]):
            color = severity_color.get(x.seriousness, AMBER)
            detail = (f"{x.frequency}. " if x.frequency else "") + x.context
            card(slide, 1.7 + row * 1.75, f"{_clip(x.event, 120)} ({x.seriousness})", _clip(detail, 260),
                 f"Evidence: {_cites(x.citations)}", color, height=1.6)

    # 8. Conflicts
    if s.conflicts:
        slide = page("Where the evidence disagrees", f"{len(s.conflicts)} adjudicated conflict(s)")
        lines = []
        for c in s.conflicts[:4]:
            lines += [f"{c.topic} ({c.severity.replace('_', ' ')})",
                      f"   A: {_clip(c.position_a, 140)} {_cites(c.citations_a)}",
                      f"   B: {_clip(c.position_b, 140)} {_cites(c.citations_b)}",
                      f"   Assessment: {_clip(c.assessment, 200)}"]
        text(slide, 0.6, 1.7, 12.1, 5.2, lines, size=12, color=INK)

    # 9. Gaps
    if s.gaps:
        slide = page("Research gaps", "what the evidence does not answer yet")
        lines = []
        for g in s.gaps[:5]:
            lines += [f"● {g.question} ({g.priority})", f"   Would answer it: {_clip(g.what_would_answer_it, 200)}"]
        text(slide, 0.6, 1.7, 12.1, 5.2, lines, size=13, color=INK)

    # 10. Sources
    for start in range(0, min(len(report.records), 36), 18):
        chunk = report.records[start:start + 18]
        slide = page("Sources" + (f" ({start // 18 + 1})" if len(report.records) > 18 else ""),
                     f"{len(report.records)} retrieved records - full list in the document export")
        lines = [f"[{r.sid}] {_clip(r.title, 95)} - {r.citation} · {DESIGN_LABELS.get(r.design, r.design)}"
                 for r in chunk]
        text(slide, 0.6, 1.6, 12.1, 5.4, lines, size=10, color=INK)

    # 11. Method and limits
    slide = page("How this was produced, and its limits")
    text(slide, 0.6, 1.6, 12.1, 5.4, [
        f"Sources: {', '.join(m['totals']['sources_used'])}.",
        "Every record was appraised individually; indices and charts were computed arithmetically, "
        "not generated by a model; every citation was verified against the retrieved corpus.",
        "Retrieval is a sample, not a systematic review - absence from the corpus is not evidence of absence.",
        "Appraisal reads abstracts and registry records, not full texts. Publication bias is not corrected for.",
        "Adverse-event report counts are reporting volume, not incidence. Regulatory coverage is US FDA only.",
        "Research intelligence for qualified professionals - not medical advice or a treatment recommendation.",
    ], size=14, color=INK)

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #

def to_docx(report: Report) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Pt, RGBColor

    s, m = report.synthesis, report.metrics
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def para(value: str, *, bold: bool = False, italic: bool = False, color: tuple | None = None, size: int | None = None):
        p = doc.add_paragraph()
        run = p.add_run(value)
        run.bold, run.italic = bold, italic
        if color:
            run.font.color.rgb = RGBColor(*color)
        if size:
            run.font.size = Pt(size)
        return p

    def table(headers: list[str], rows: list[list[Any]]) -> None:
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for cell, header in zip(t.rows[0].cells, headers):
            cell.text = header
            cell.paragraphs[0].runs[0].bold = True
        for row in rows:
            cells = t.add_row().cells
            for cell, value in zip(cells, row):
                cell.text = str(value)
        doc.add_paragraph()

    doc.add_heading("DrugScope research report", level=0)
    para(report.query, bold=True, size=14)
    para(f"Interpreted as: {report.plan.interpretation}", italic=True, color=MUTED)
    para(f"Generated {report.generated_at} · {report.depth} depth · {len(report.records)} sources",
         color=MUTED, size=9)

    doc.add_heading("Bottom line", level=1)
    para(s.summary.verdict, bold=True, size=13)
    para(f"Confidence: {s.summary.confidence}. {s.summary.confidence_rationale}")
    table(["Index", "Value"], [
        ["Evidence Quality Index", f"{m['quality']['score']}/100"],
        ["Consensus Index", f"{m['consensus']['score']}/100"],
        ["Development maturity", f"{s.summary.maturity}/100"],
        ["Sources analysed", m["totals"]["records"]],
        ["Structured claims", m["totals"]["claims"]],
        ["Evidence span", m["year_span"]],
    ])

    doc.add_heading("Executive summary", level=1)
    for paragraph in [p for p in s.summary.narrative.split("\n") if p.strip()]:
        para(paragraph)
    if s.summary.watch_items:
        doc.add_heading("What to watch", level=2)
        for item in s.summary.watch_items:
            doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("Key findings", level=1)
    numbers = (m.get("number_check") or {}).get("findings") or {}
    for i, finding in enumerate(s.key_findings, 1):
        doc.add_heading(f"{i}. {finding.headline}", level=2)
        para(f"Strength: {finding.strength} · Evidence: {_cites(finding.citations)}", italic=True, color=MUTED)
        para(finding.detail)
        para(f"So what: {finding.so_what}", bold=True)
        check = numbers.get(i) or numbers.get(str(i))
        if check and check.get("unverified"):
            para("Figures not found verbatim in the cited sources (may be derived): "
                 + ", ".join(check["unverified"]), italic=True, color=AMBER, size=9)

    if s.comparison.entities and s.comparison.rows:
        doc.add_heading("Comparison", level=1)
        rows = []
        for row in s.comparison.rows:
            by_entity = {c.entity: c for c in row.cells}
            rows.append([row.criterion] + [
                f"{by_entity[e].value} {_cites(by_entity[e].citations)}" if e in by_entity else "Not established"
                for e in s.comparison.entities] + [row.verdict or "-"])
        table(["Criterion", *s.comparison.entities, "Verdict"], rows)
        para(f"Bottom line: {s.comparison.bottom_line}", bold=True)

    if s.timeline:
        doc.add_heading("Development timeline", level=1)
        table(["Date", "Event", "Category", "Sources"],
              [[e.date, f"{e.label} - {e.detail}", e.category, _cites(e.citations)]
               for e in sorted(s.timeline, key=lambda e: e.date or "9999")])

    doc.add_heading("Conflicting evidence", level=1)
    if s.conflicts:
        for c in s.conflicts:
            doc.add_heading(f"{c.topic} ({c.severity.replace('_', ' ')})", level=2)
            para(f"Position A: {c.position_a} {_cites(c.citations_a)}")
            para(f"Position B: {c.position_b} {_cites(c.citations_b)}")
            para(f"Assessment: {c.assessment}", bold=True)
            para(f"Likely cause of divergence: {c.likely_explanation}", italic=True)
    else:
        para("No direct contradictions were identified. Internal consistency is not the same as correctness.")

    if s.gaps:
        doc.add_heading("Research gaps", level=1)
        for i, g in enumerate(s.gaps, 1):
            doc.add_heading(f"{i}. {g.question} ({g.priority})", level=2)
            para(f"Why it matters: {g.why_it_matters}")
            para(f"Missing from the corpus: {g.evidence_absent}")
            para(f"What would answer it: {g.what_would_answer_it}")

    if s.safety_signals:
        doc.add_heading("Safety signals", level=1)
        table(["Event", "Seriousness", "Frequency", "Context", "Sources"],
              [[x.event, x.seriousness, x.frequency or "-", x.context, _cites(x.citations)] for x in s.safety_signals])

    doc.add_heading("Regulatory status", level=1)
    para(s.regulatory_status)
    doc.add_heading("Mechanism", level=1)
    para(s.mechanism)
    briefing = m.get("recent_developments", "")
    if briefing and not briefing.startswith("("):
        doc.add_heading("Recent developments (open web)", level=1)
        para(briefing)

    doc.add_heading("Sources", level=1)
    for r in report.records:
        p = doc.add_paragraph()
        p.add_run(f"[{r.sid}] ").bold = True
        p.add_run(f"{r.title}. {r.citation}. {DESIGN_LABELS.get(r.design, r.design)}, {r.tier}. ")
        if r.url:
            p.add_run(r.url).italic = True

    doc.add_heading("Methodology and limitations", level=1)
    para(f"Evidence was retrieved by direct query to {', '.join(m['totals']['sources_used'])}. Each record was "
         "appraised individually; counts, indices and charts were computed arithmetically, not generated by a "
         "language model; every citation was verified against the retrieved corpus.")
    para("Evidence tiers: " + "; ".join(f"{t} - {meaning}" for t, meaning in TIER_MEANING.items()) + ".")
    for limit in (
        "Retrieval is a sample, not a systematic review; absence from this corpus is not evidence of absence.",
        "Appraisal is based on abstracts and registry records, not full texts.",
        "Publication bias is not corrected for.",
        "Adverse-event report counts are reporting volume, not incidence, and do not establish causation.",
        "Regulatory coverage is US FDA only.",
        "Research intelligence for qualified professionals - not medical advice or a treatment recommendation.",
    ):
        doc.add_paragraph(limit, style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
